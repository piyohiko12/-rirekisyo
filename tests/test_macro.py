"""履歴書ブックのマクロ（VBA）が壊れていないかを確かめる。

VBAはExcelの外では動かせないので、ここでは
・構文の骨組み（ブロックの対応・引用符・識別子）
・マクロが当てにしているシートやセルが本当にブックにあるか
を確かめる。手で貼り付けても壊れないこと（二重引用符の扱い）も見る。
"""

import re

import openpyxl
import pytest

import make_excel_book as m

BOOK = "履歴書作成.xlsm"


@pytest.fixture(scope="module")
def book():
    """履歴書ブックは大きいので、1回だけ読んで使い回す。"""
    return openpyxl.load_workbook(BOOK)


def logical_lines(lines):
    """行継続( _ )でつながった行を1行にまとめる。"""
    out, buf, start = [], "", 0
    for n, raw in enumerate(lines, 1):
        line = raw.rstrip()
        if buf == "":
            start = n
        if line.endswith(" _"):
            buf += line[:-1]
            continue
        out.append((start, buf + line))
        buf = ""
    assert buf == "", "行継続( _ )で終わったまま終わっている"
    return out


OPENERS = (
    (re.compile(r"^\s*(Public |Private )?(Sub|Function)\s"), "proc"),
    (re.compile(r"^\s*(For\s+Each\b|For\b)"), "For"),
    (re.compile(r"^\s*Do\b"), "Do"),
    (re.compile(r"^\s*With\b"), "With"),
    (re.compile(r"^\s*Select\s+Case\b"), "Select"),
)
CLOSERS = (
    (re.compile(r"^End (Sub|Function)\b"), "proc"),
    (re.compile(r"^Next\b"), "For"),
    (re.compile(r"^Loop\b"), "Do"),
    (re.compile(r"^End With\b"), "With"),
    (re.compile(r"^End Select\b"), "Select"),
    (re.compile(r"^End If\b"), "If"),
)


def check_vba(lines):
    """VBAの骨組みを検査して、見つかった問題を返す。"""
    errs, stack = [], []
    for n, text in logical_lines(lines):
        if text.count('"') % 2:
            errs.append(f"{n}: 二重引用符の数が奇数")
        code = re.sub(r'"[^"]*"', '""', text)
        code = re.sub(r"'.*$", "", code)
        if not code.strip():
            continue
        if code.count("(") != code.count(")"):
            errs.append(f"{n}: 括弧が合わない -> {text.strip()[:50]}")
        for name in re.findall(r"\b(?:Dim|Const|Sub|Function)\s+([^\s(,]+)", code):
            if name[:1].isdigit():
                errs.append(f"{n}: 数字で始まる識別子 {name}")
        s = code.strip()
        for pat, kind in CLOSERS:
            if pat.match(s):
                if not stack or stack[-1] != kind:
                    errs.append(f"{n}: 対応しない {s[:20]}")
                else:
                    stack.pop()
                break
        else:
            if re.match(r"^If\b.*\bThen\s*$", s):
                stack.append("If")
                continue
            for pat, kind in OPENERS:
                if pat.match(s):
                    stack.append(kind)
                    break
    if stack:
        errs.append(f"閉じられていないブロック: {stack}")
    return errs


def test_macro_code_structure():
    assert check_vba(m.macro_code()) == []


def test_thisworkbook_code_structure():
    assert check_vba(m.THISWORKBOOK_CODE) == []


def test_bas_file_is_importable(tmp_path):
    """.bas は Shift_JIS・CRLF で、モジュール名の宣言から始まること。"""
    path = tmp_path / m.BAS_FILE_NAME
    m.write_bas(path)
    raw = path.read_bytes()
    text = raw.decode("cp932")            # 読めなければここで落ちる
    assert text.splitlines()[0] == f'Attribute VB_Name = "{m.BAS_MODULE_NAME}"'
    assert raw.count(b"\r\n") == text.count("\n")


def test_measure_sheet_exists_and_is_hidden(book):
    """マクロが高さを測るのに使う作業シートが、非表示で入っていること。"""
    wb = book
    assert m.MEASURE in wb.sheetnames
    assert wb[m.MEASURE].sheet_state == "hidden"
    assert wb[m.MEASURE].print_area in (None, "")


def test_macro_points_at_sheets_that_exist(book):
    """VBAの中の Const で書いたシート名が、実際のブックにあること。"""
    wb = book
    code = "\n".join(m.macro_code())
    names = re.findall(r'Const SHEET_\w+ As String = "([^"]+)"', code)
    assert names, "シート名の定数が見つからない"
    missing = [n for n in names if n not in wb.sheetnames]
    assert missing == [], missing


@pytest.mark.parametrize("cell_name", [
    "MOTIVATION_CELL", "ACTIVITIES_CELL", "REMARKS_CELL",
    "NAME_CELL", "NAME_KANA_CELL", "ADDR_CELL", "ADDR_KANA_CELL",
    "CONTACT_CELL", "CONTACT_KANA_CELL", "SCHOOL_CELL",
])
def test_macro_targets_are_merged_cells(book, cell_name):
    """マクロが大きさを変える欄は、1つの結合セルになっていること。

    結合されていないと、文章の先頭セルしか読めず判定を誤る。
    """
    ws = book[m.FORM_SHEET]
    ranges = {str(r) for r in ws.merged_cells.ranges}
    assert getattr(m, cell_name) in ranges


def test_macro_sheet_backup_has_no_double_quotes(book):
    """「マクロ」シートのC列（貼り付け用の予備）に " が残っていないこと。

    Excelはコピーのときに " を二重にしてしまうため、全角の ” にしてある。
    """
    ws = book[m.MACRO_SHEET]
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row=row, column=3).value
        assert '"' not in str(value or "")
