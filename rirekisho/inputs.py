"""入力シート（Excel / Googleスプレッドシート）の定義・生成・読み取り。

**1行＝1生徒** の名列順で、クラス全員分をまとめて入力する形にしている
（調査書サポートの「入力」シートと同じ並べ方）。

    1行目: タイトル
    2行目: 大見出し（基本情報 / 現住所 / 資格等 …）
    3行目: 列見出し
    4行目: キー（非表示。読み取りに使うので消さない）
    5行目以降: 生徒1人につき1行（既定40人分）

読み取りはキー行だけを見るので、列を動かしても、Googleスプレッドシートで
.xlsx や .csv に書き出しても同じように動く。
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path

TITLE_ROW = 1
GROUP_ROW = 2
HEADER_ROW = 3
KEY_ROW = 4
FIRST_DATA_ROW = 5
ROSTER_ROWS = 40          # 既定の名簿の行数
SHEET_NAME = "入力"
SETTINGS_SHEET = "設定"
MASTER_SHEET = "資格マスタ"
PASTE_SHEET = "資格取込"
FIELDS_SHEET = "反映項目"
PASTE_FIRST_ROW = 6  # 貼り付けを始める行

MAX_LICENSE_SLOTS = 6     # 資格の入力枠（用紙は5行だが多めに入力できる）
MAX_JOB_SLOTS = 2         # 職歴の入力枠


@dataclass
class Column:
    label: str
    key: str
    width: float = 14.0
    group: str = ""
    note: str = ""
    kind: str = "text"  # text / date / int


def columns() -> list[Column]:
    """入力シートの列（左から順）。"""
    cols: list[Column] = [
        Column("No", "meta.no", 4.5, "生徒", kind="int"),
        Column("組", "meta.class", 9.0, "生徒", "例: 3年2組"),
        Column("出席番号", "meta.number", 7.0, "生徒", kind="int"),
        Column("氏名", "name", 14.0, "生徒", "例: 佐野 太郎"),
        Column("ふりがな（氏名）", "name_kana", 16.0, "生徒", "例: さの たろう"),
        Column("生年月日", "birth", 12.5, "生徒", "2008-05-12 / 平成20年5月12日", kind="date"),
        Column("郵便番号", "zip", 10.0, "現住所", "5980001 でも可"),
        Column("住所", "address", 34.0, "現住所"),
        Column("ふりがな（住所）", "address_kana", 26.0, "現住所"),
        Column("連絡先 郵便番号", "contact_zip", 11.0, "連絡先", "空欄なら「同上」"),
        Column("連絡先 住所", "contact_address", 26.0, "連絡先", "空欄なら「同上」"),
        Column("連絡先 ふりがな", "contact_kana", 20.0, "連絡先"),
    ]
    for i in range(1, MAX_LICENSE_SLOTS + 1):
        cols.append(Column(f"資格{i} 名称", f"license.{i}.name", 22.0, "資格等"))
        cols.append(
            Column(f"資格{i} 取得年月", f"license.{i}.ym", 12.0, "資格等",
                   "2024-06 / 令和6年6月", kind="date")
        )
    cols += [
        Column("校内外の諸活動", "activities", 34.0, "校内外の諸活動", "セル内改行で箇条書き"),
        Column("志望の動機", "motivation", 40.0, "志望の動機など",
               "この欄だけでも可。見出しは付けずにそのまま印字します"),
        Column("希望の職種", "desired_job", 18.0, "志望の動機など", "任意。志望の動機に続けて印字します"),
        Column("アピールポイント", "appeal", 34.0, "志望の動機など", "任意。最後に続けて印字します"),
    ]
    for i in range(1, MAX_JOB_SLOTS + 1):
        cols.append(Column(f"職歴{i} 年月", f"job.{i}.ym", 11.0, "職歴", kind="date"))
        cols.append(Column(f"職歴{i} 内容", f"job.{i}.text", 26.0, "職歴"))
    cols.append(Column("備考", "remarks", 26.0, "備考"))
    return cols


SETTINGS = [
    ("基準日（満○歳の計算日）", "as_of", "2026-09-01",
     "用紙に印字された「令和8年9月1日現在」に合わせています"),
    ("学級（ファイル名に使う）", "class_label", "",
     "空欄なら各行の「組」を使います"),
    ("資格の並び順", "license_order", "取得年月順",
     "「取得年月順」または「入力順」"),
]

# 資格マスタの初期値（入力・取り込みでの名称 → 履歴書に印字する正式名称）。
# 学校で使っている資格名の一覧。シート上で追加・修正できる。
DEFAULT_MASTER = [
    ("実用英語検定5級", "実用英語検定5級"),
    ("実用英語検定4級", "実用英語検定4級"),
    ("実用英語検定3級", "実用英語検定3級"),
    ("実用英語検定準２級", "実用英語検定準２級"),
    ("実用英語検定2級", "実用英語検定2級"),
    ("実用英語検定1級", "実用英語検定1級"),
    ("日本漢字能力検定4級", "日本漢字能力検定4級"),
    ("日本漢字能力検定3級", "日本漢字能力検定3級"),
    ("日本漢字能力検定準2級", "日本漢字能力検定準2級"),
    ("日本漢字能力検定2級", "日本漢字能力検定2級"),
    ("計算技術検定1級", "計算技術検定1級"),
    ("計算技術検定2級", "計算技術検定2級"),
    ("計算技術検定3級", "計算技術検定3級"),
    ("計算技術検定4級", "計算技術検定4級"),
    ("情報技術検定1級", "情報技術検定1級"),
    ("情報技術検定2級", "情報技術検定2級"),
    ("情報技術検定3級", "情報技術検定3級"),
    ("初級CAD検定", "初級CAD検定"),
    ("基礎製図検定", "基礎製図検定"),
    ("機械製図検定", "機械製図検定"),
    ("色彩検定３級", "色彩検定３級"),
    ("色彩検定２級", "色彩検定２級"),
    ("色彩検定UC級", "色彩検定UC級"),
    ("QC検定4級", "QC検定4級"),
    ("QC検定3級", "QC検定3級"),
    ("日本語ワープロ検定2級", "日本語ワープロ検定2級"),
    ("日本語ワープロ検定準2級", "日本語ワープロ検定準2級"),
    ("日本語ワープロ検定3級", "日本語ワープロ検定3級"),
    ("日本語ワープロ検定4級", "日本語ワープロ検定4級"),
    ("情報処理技能検定(表計算)4級", "情報処理技能検定(表計算)4級"),
    ("文書デザイン検定4級", "文書デザイン検定4級"),
    ("プレゼンテーション作成検定3級", "プレゼンテーション作成検定3級"),
    ("プレゼンテーション作成検定4級", "プレゼンテーション作成検定4級"),
    ("パソコンスピード認定1級", "パソコンスピード認定1級"),
    ("パソコンスピード認定2級", "パソコンスピード認定2級"),
    ("パソコンスピード認定3級", "パソコンスピード認定3級"),
    ("パソコンスピード認定4級", "パソコンスピード認定4級"),
    ("パソコンスピード認定5級", "パソコンスピード認定5級"),
    ("電気工事士第一種", "電気工事士第一種"),
    ("電気工事士第二種", "電気工事士第二種"),
    ("危険物取扱者乙1", "危険物取扱者乙1"),
    ("危険物取扱者乙2", "危険物取扱者乙2"),
    ("危険物取扱者乙3", "危険物取扱者乙3"),
    ("危険物取扱者乙4", "危険物取扱者乙4"),
    ("危険物取扱者乙5", "危険物取扱者乙5"),
    ("危険物取扱者乙6", "危険物取扱者乙6"),
    ("危険物取扱者丙", "危険物取扱者丙"),
    ("ガス溶接技能講習修了", "ガス溶接技能講習修了"),
    ("ボイラー取扱技能講習修了", "ボイラー取扱技能講習修了"),
    ("ア－ク溶接安全衛生教育修了", "ア－ク溶接安全衛生教育修了"),
    ("小型フォークリフト運転特別教育講習修了", "小型フォークリフト運転特別教育講習修了"),
    ("二級ボイラー技士2級", "二級ボイラー技士2級"),
    ("技能検定　機械加工（普通旋盤作業）2級", "技能検定　機械加工（普通旋盤作業）2級"),
    ("技能検定　機械加工（普通旋盤作業）3級", "技能検定　機械加工（普通旋盤作業）3級"),
    ("技能検定　機械検査（機械検査作業）2級", "技能検定　機械検査（機械検査作業）2級"),
    ("技能検定　機械検査（機械検査作業）3級", "技能検定　機械検査（機械検査作業）3級"),
    ("アーク溶接技能者適格性証明書Ａ－２Ｆ", "アーク溶接技能者適格性証明書Ａ－２Ｆ"),
    ("アーク溶接技能者適格性証明書Ａ－２V", "アーク溶接技能者適格性証明書Ａ－２V"),
    ("アーク溶接技能者適格性証明書Ｎ－２Ｆ", "アーク溶接技能者適格性証明書Ｎ－２Ｆ"),
    ("アーク溶接技能者適格性証明書Ａ－２H", "アーク溶接技能者適格性証明書Ａ－２H"),
    ("アーク溶接技能者適格性証明書SＮ－２Ｆ", "アーク溶接技能者適格性証明書SＮ－２Ｆ"),
    ("アーク溶接技能者適格性証明書T－1Ｆ", "アーク溶接技能者適格性証明書T－1Ｆ"),
    ("アーク溶接技能者適格性証明書SA－2Ｆ", "アーク溶接技能者適格性証明書SA－2Ｆ"),
    ("第二級デジタル通信", "第二級デジタル通信"),
    ("ジュニアマイスターゴールド", "ジュニアマイスターゴールド"),
    ("ジュニアマイスターシルバー", "ジュニアマイスターシルバー"),
    ("ジュニアマイスターブロンズ", "ジュニアマイスターブロンズ"),
    ("フォークリフト運転技能講習", "フォークリフト運転技能講習"),
    ("玉掛け技能講習", "玉掛け技能講習"),
    ("硬筆書写技能検定4級", "硬筆書写技能検定4級"),
]


# PDFに反映する項目（キー, 画面やシートに出す名前, 既定で反映するか）
OUTPUT_FIELDS = [
    ("name", "氏名", True),
    ("name_kana", "ふりがな（氏名）", True),
    ("birth", "生年月日・満○歳", True),
    ("zip", "郵便番号", True),
    ("address", "住所", True),
    ("address_kana", "ふりがな（住所）", True),
    ("contact", "連絡先（「同上」を含む）", True),
    ("licenses", "資格等", True),
    ("activities", "校内外の諸活動", True),
    ("motivation", "志望の動機", True),
    ("desired_job", "希望の職種", True),
    ("appeal", "アピールポイント", True),
    ("jobs", "職歴", True),
    ("remarks", "備考", True),
]

ON_MARKS = {"○", "◯", "〇", "o", "yes", "true", "1", "はい", "on", "●", "✓", "レ"}
OFF_MARKS = {"×", "x", "no", "false", "0", "いいえ", "off", "-", "―", "ー"}


def default_fields() -> dict[str, bool]:
    """既定の反映設定（すべて反映）。"""
    return {key: default for key, _label, default in OUTPUT_FIELDS}


def parse_mark(value, default: bool = True) -> bool:
    """「○／×」などの記号を True / False にする。"""
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if not text:
        return default
    if text in ON_MARKS:
        return True
    if text in OFF_MARKS:
        return False
    return default


@dataclass
class Student:
    """入力シートの1行（生徒1人分）。"""

    row: int
    no: int | None
    class_name: str
    number: str
    values: dict[str, object] = field(default_factory=dict)

    @property
    def name(self) -> str:
        return str(self.values.get("name") or "").strip()

    @property
    def is_empty(self) -> bool:
        return not any(
            str(v).strip() for k, v in self.values.items() if not k.startswith("meta.") and v
        )

    def file_stem(self, class_label: str = "") -> str:
        """出力するPDFのファイル名（拡張子なし）。"""
        head = self.number or (str(self.no) if self.no else "")
        parts = [class_label or self.class_name, head.zfill(2) if head.isdigit() else head]
        parts = [p for p in parts if p]
        name = self.name.replace(" ", "").replace("　", "") or "無名"
        for bad in '\\/:*?"<>|':
            name = name.replace(bad, "_")
        return "_".join([*parts, name])


# ------------------------------------------------------------------ 生成
def write_template(
    path: str | Path,
    *,
    students: list[dict[str, object]] | None = None,
    rows: int = ROSTER_ROWS,
    settings: dict[str, object] | None = None,
    paste_lines: list[str] | None = None,
    fields: dict[str, bool] | None = None,
) -> Path:
    """名列順の入力用 .xlsx を生成する。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    path = Path(path)
    cols = columns()
    students = students or []

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    input_fill = PatternFill("solid", fgColor="FFFDE7")
    meta_fill = PatternFill("solid", fgColor="EFEFEF")
    header_fill = PatternFill("solid", fgColor="E8EEF4")
    group_fill = PatternFill("solid", fgColor="D6E2EF")
    small = Font(size=9, color="666666")

    ws.cell(row=TITLE_ROW, column=1, value="履歴書 入力シート（1行＝1生徒・名列順）").font = Font(
        size=13, bold=True
    )
    ws.cell(
        row=TITLE_ROW,
        column=4,
        value="黄色いセルに入力して保存 → python build_pdf.py で全員分のPDFを作ります。"
        "4行目（非表示のキー行）は消さないでください。",
    ).font = small

    # 2行目: 大見出し（同じ見出しが続く範囲を結合する）
    start = 1
    for i, col in enumerate(cols, start=1):
        last = i == len(cols)
        if last or cols[i].group != col.group:
            cell = ws.cell(row=GROUP_ROW, column=start, value=col.group)
            cell.font = Font(size=10, bold=True)
            cell.alignment = Alignment(horizontal="center")
            if i > start:
                ws.merge_cells(start_row=GROUP_ROW, start_column=start, end_row=GROUP_ROW, end_column=i)
            for c in range(start, i + 1):
                ws.cell(row=GROUP_ROW, column=c).fill = group_fill
            start = i + 1

    # 3行目: 列見出し / 4行目: キー
    for i, col in enumerate(cols, start=1):
        head = ws.cell(row=HEADER_ROW, column=i, value=col.label)
        head.fill, head.border = header_fill, border
        head.font = Font(size=10, bold=True)
        head.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if col.note:
            head.comment = _comment(col.note)
        ws.cell(row=KEY_ROW, column=i, value=col.key)
        ws.column_dimensions[get_column_letter(i)].width = col.width
    ws.row_dimensions[KEY_ROW].hidden = True
    ws.row_dimensions[HEADER_ROW].height = 30

    # 5行目以降: 生徒
    for r in range(rows):
        row = FIRST_DATA_ROW + r
        data = students[r] if r < len(students) else {}
        for i, col in enumerate(cols, start=1):
            value = data.get(col.key)
            if col.key == "meta.no" and value is None:
                value = r + 1
            cell = ws.cell(row=row, column=i, value=value)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = meta_fill if col.key.startswith("meta.") else input_fill
            if col.kind == "date":
                cell.number_format = "yyyy/mm/dd"
        ws.row_dimensions[row].height = 22

    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=5)  # No〜氏名を固定
    ws.auto_filter.ref = (
        f"A{HEADER_ROW}:{get_column_letter(len(cols))}{FIRST_DATA_ROW + rows - 1}"
    )

    # 設定シート
    st = wb.create_sheet(SETTINGS_SHEET)
    st.cell(row=1, column=1, value="設定（全員に共通）").font = Font(size=12, bold=True)
    settings = settings or {}
    for i, (label, key, default, note) in enumerate(SETTINGS, start=3):
        st.cell(row=i, column=1, value=label)
        cell = st.cell(row=i, column=2, value=settings.get(key, default))
        cell.fill, cell.border = input_fill, border
        st.cell(row=i, column=3, value=note).font = small
        st.cell(row=i, column=4, value=key)
    st.column_dimensions["A"].width = 26
    st.column_dimensions["B"].width = 20
    st.column_dimensions["C"].width = 52
    st.column_dimensions["D"].hidden = True

    # 反映項目（PDFに出す項目を○×で選ぶ）
    fields = {**default_fields(), **(fields or {})}
    fs = wb.create_sheet(FIELDS_SHEET)
    fs.cell(row=1, column=1, value="反映項目（PDFに出す項目を選びます）").font = Font(size=12, bold=True)
    fs.cell(
        row=2,
        column=1,
        value="B列を「○」にした項目だけをPDFに書き込みます。「×」にすると、入力してあっても"
        "その欄は空欄のまま印刷されます（入力シートの値は消えません）。",
    ).font = small
    for col, label in ((1, "項目"), (2, "反映する（○／×）")):
        cell = fs.cell(row=3, column=col, value=label)
        cell.fill, cell.border, cell.font = header_fill, border, Font(size=10, bold=True)
    marks = DataValidation(type="list", formula1='"○,×"', allow_blank=True)
    fs.add_data_validation(marks)
    for i, (key, label, _default) in enumerate(OUTPUT_FIELDS, start=4):
        fs.cell(row=i, column=1, value=label).border = border
        cell = fs.cell(row=i, column=2, value="○" if fields.get(key, True) else "×")
        cell.fill, cell.border = input_fill, border
        cell.alignment = Alignment(horizontal="center")
        marks.add(cell)
        fs.cell(row=i, column=3, value=key)
    fs.column_dimensions["A"].width = 26
    fs.column_dimensions["B"].width = 18
    fs.column_dimensions["C"].hidden = True
    fs.freeze_panes = "A4"

    # 資格マスタ
    ms = wb.create_sheet(MASTER_SHEET)
    ms.cell(row=1, column=1, value="資格マスタ（入力・取り込みの名称 → 履歴書に印字する正式名称）").font = Font(
        size=12, bold=True
    )
    ms.cell(
        row=2,
        column=1,
        value="ここに無い資格名は、入力したままの名称で印字します。行はいくらでも追加できます。"
        "空白や全角半角の違いは自動で吸収します。",
    ).font = small
    for col, label in ((1, "入力・取り込みでの名称"), (2, "履歴書での正式名称"), (3, "メモ")):
        cell = ms.cell(row=3, column=col, value=label)
        cell.fill, cell.border, cell.font = header_fill, border, Font(size=10, bold=True)
    for i, (src, dest) in enumerate(DEFAULT_MASTER, start=4):
        for col, value in ((1, src), (2, dest)):
            cell = ms.cell(row=i, column=col, value=value)
            cell.fill, cell.border = input_fill, border
        ms.cell(row=i, column=3).border = border
    ms.column_dimensions["A"].width = 34
    ms.column_dimensions["B"].width = 34
    ms.column_dimensions["C"].width = 30
    ms.freeze_panes = "A4"

    # 資格取込（貼り付け用）
    ps = wb.create_sheet(PASTE_SHEET)
    ps.cell(row=1, column=1, value="資格取込（1行＝1件で貼り付け）").font = Font(size=12, bold=True)
    ps.cell(
        row=2,
        column=1,
        value=f"A{PASTE_FIRST_ROW}以降に、資格取得の一覧を1行1件で貼り付けてください。"
        "build_pdf.py の実行時に名簿と照合して、各生徒の資格欄に自動で追加します。",
    ).font = small
    ps.cell(
        row=3,
        column=1,
        value="書式: 3-2-15〔空白またはTAB〕山田太郎 基礎製図検定 令和6年7月10日"
        "　…学年-組-出席番号 → 氏名 → 資格名 → 取得日 の順。"
        "先頭の番号がない場合は氏名で照合します。資格名は「資格マスタ」で正式名称に変換します。",
    ).font = small
    ps.cell(
        row=4,
        column=1,
        value="結果は 出力/資格取込ログ.csv に書き出します（このシートには書き戻しません）。",
    ).font = small
    head = ps.cell(row=PASTE_FIRST_ROW - 1, column=1, value="貼付原文")
    head.fill, head.border, head.font = header_fill, border, Font(size=10, bold=True)
    paste_lines = paste_lines or []
    for i, r in enumerate(range(PASTE_FIRST_ROW, PASTE_FIRST_ROW + 200)):
        cell = ps.cell(row=r, column=1, value=paste_lines[i] if i < len(paste_lines) else None)
        cell.fill, cell.border = input_fill, border
    ps.column_dimensions["A"].width = 80
    ps.freeze_panes = f"A{PASTE_FIRST_ROW}"

    guide = wb.create_sheet("使い方")
    for i, line in enumerate(USAGE_LINES, start=1):
        cell = guide.cell(row=i, column=1, value=line)
        cell.alignment = Alignment(vertical="top")
        if line.startswith("■"):
            cell.font = Font(bold=True)
    guide.column_dimensions["A"].width = 100

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _comment(text: str):
    from openpyxl.comments import Comment

    comment = Comment(text, "履歴書システム")
    comment.width, comment.height = 240, 60
    return comment


USAGE_LINES = [
    "■ このファイルについて",
    "「入力」シートは 1行＝1生徒 の名列順です。クラス全員分をまとめて入力できます。",
    "入力して保存したあと python build_pdf.py を実行すると、生徒ごとの履歴書PDFが 出力/ にできます。",
    "",
    "■ 手順",
    "1. 「入力」シートの黄色いセルに、名列順で入力する（氏名が空の行はPDFを作りません）。",
    "2. 「設定」シートの基準日を確認する（満○歳の計算に使います）。",
    "3. このファイルを保存して閉じる（開いたままでも作れます）。",
    "4. 「履歴書PDF作成」を起動して、青い【PDF作成】ボタンを押す。",
    "   Windows … 履歴書PDF作成.bat をダブルクリック",
    "   macOS  … 履歴書PDF作成.command をダブルクリック",
    "   ※ 画面では「全員／No.指定／氏名で絞り込み」「1つのPDFにまとめる」「自動更新」も選べます。",
    "",
    "■ コマンドで使う場合（画面を使わないとき）",
    "   python build_pdf.py            … 全員分のPDFを作る",
    "   python build_pdf.py --no 3     … No.3 の生徒だけ",
    "   python build_pdf.py --merge 3年2組_履歴書.pdf … 全員を1つのPDFにまとめる（印刷用）",
    "   python watch.py                … 保存するたび自動で作り直す",
    "",
    "■ PDFに出す項目を選ぶ",
    "「反映項目」シートで、項目ごとに ○（出す）／×（出さない）を選べます。",
    "×にした欄は、入力してあってもPDFでは空欄のままになります（入力シートの値は消えません）。",
    "画面（履歴書PDF作成）のチェックボックスでも、その場で切り替えられます。",
    "",
    "■ 自動で入る項目",
    "・満○歳 … 生年月日と基準日から自動計算します。",
    "・元号（昭和／平成／令和） … 生年月日から自動判定します。",
    "・連絡先 … 空欄なら「同上」と印字します。",
    "・郵便番号 … 7桁の数字だけでも 123-4567 の形に整えます。",
    "・資格の取得年月 … 西暦で入力しても和暦（例: 令和6年6月）で印字します。",
    "",
    "■ 注意",
    "・4行目は非表示のキー行です。読み取りに使うので、消したり並べ替えたりしないでください。",
    "・生徒を増やすときは、5行目以降の行をコピーして貼り付けてください。",
    "・列を増やすときは、4行目に同じ書き方のキー（license.7.name など）も入れてください。",
    "・写真は用紙の「写真をはる位置」に貼ってください（PDFには合成しません）。",
]


# ------------------------------------------------------------------ 読み取り
@dataclass
class InputData:
    """入力シートから読み取った内容一式。"""

    students: list[Student] = field(default_factory=list)
    settings: dict[str, object] = field(default_factory=dict)
    fields: dict[str, bool] = field(default_factory=default_fields)
    master_pairs: list[tuple[object, object]] = field(default_factory=list)
    paste_lines: list[str] = field(default_factory=list)


def read_students(path: str | Path) -> InputData:
    """入力シート（.xlsx / .csv）を読み取る。

    氏名などが何も入っていない行は読み飛ばす。
    以前の「1人1枚」形式のシートも、生徒1人として読み込む。
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"入力ファイルが見つかりません: {path}")
    if path.suffix.lower() == ".csv":
        rows = _csv_rows(path)
        return InputData(students=_students_from_rows(rows, source=path))
    return _read_xlsx(path)


def _read_xlsx(path: Path) -> InputData:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]
    rows = [[c.value for c in row] for row in ws.iter_rows()]

    data = InputData()
    if SETTINGS_SHEET in wb.sheetnames:
        for row in wb[SETTINGS_SHEET].iter_rows(values_only=True):
            if len(row) >= 4 and isinstance(row[3], str) and row[3].strip():
                data.settings[row[3].strip()] = row[1]

    if FIELDS_SHEET in wb.sheetnames:
        for row in wb[FIELDS_SHEET].iter_rows(values_only=True):
            if len(row) >= 3 and isinstance(row[2], str) and row[2].strip() in data.fields:
                data.fields[row[2].strip()] = parse_mark(row[1])

    if MASTER_SHEET in wb.sheetnames:
        for i, row in enumerate(wb[MASTER_SHEET].iter_rows(values_only=True), start=1):
            if i <= 3 or len(row) < 2:  # 1〜3行目は見出し
                continue
            data.master_pairs.append((row[0], row[1]))

    if PASTE_SHEET in wb.sheetnames:
        for i, row in enumerate(wb[PASTE_SHEET].iter_rows(values_only=True), start=1):
            if i < PASTE_FIRST_ROW or not row:
                continue
            # 1セルに1件が基本だが、列が分かれて貼られても拾えるようにつなぐ
            line = " ".join(_as_paste_text(c) for c in row if c is not None).strip()
            if line:
                data.paste_lines.append(line)

    if _find_key_row(rows) is None and _looks_like_single_sheet(rows):
        data.students = [_student_from_single_sheet(rows)]
    else:
        data.students = _students_from_rows(rows, source=path)
    return data


def _as_paste_text(value) -> str:
    import datetime as _dt

    if isinstance(value, _dt.datetime):
        return value.date().isoformat()
    if isinstance(value, _dt.date):
        return value.isoformat()
    return _as_text(value)


def _csv_rows(path: Path) -> list[list[object]]:
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return [list(row) for row in csv.reader(fh)]


def _find_key_row(rows: list[list[object]]) -> int | None:
    """キー行（name などのキーが並ぶ非表示の行）の位置を返す。"""
    for i, row in enumerate(rows):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if "name" in cells and ("meta.no" in cells or "birth" in cells):
            return i
    return None


def _students_from_rows(rows: list[list[object]], *, source: Path) -> list[Student]:
    key_row = _find_key_row(rows)
    if key_row is None:
        raise ValueError(
            f"{source} からキー行が見つかりませんでした。"
            "make_xlsx.py で作った入力シートを使うか、4行目のキー行を残してください。"
        )
    keys = [str(c).strip() if c is not None else "" for c in rows[key_row]]

    students: list[Student] = []
    for offset, row in enumerate(rows[key_row + 1 :], start=key_row + 2):
        values = {key: row[i] for i, key in enumerate(keys) if key and i < len(row)}
        student = Student(
            row=offset,
            no=_as_int(values.get("meta.no")),
            class_name=_as_text(values.get("meta.class")),
            number=_as_text(values.get("meta.number")),
            values=values,
        )
        if student.is_empty:
            continue
        students.append(student)
    return students


def _as_int(value) -> int | None:
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return None


def _as_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# ---- 以前の「1人1枚」形式のシートも読めるようにしておく ------------------
_SINGLE_KEY_COL_B = 5  # E列
_SINGLE_KEY_COL_C = 6  # F列


def _looks_like_single_sheet(rows: list[list[object]]) -> bool:
    return any(
        len(row) >= _SINGLE_KEY_COL_B
        and isinstance(row[_SINGLE_KEY_COL_B - 1], str)
        and row[_SINGLE_KEY_COL_B - 1].strip() == "name"
        for row in rows
    )


def _student_from_single_sheet(rows: list[list[object]]) -> Student:
    values: dict[str, object] = {}
    for row in rows:
        for key_idx, value_idx in ((_SINGLE_KEY_COL_B - 1, 1), (_SINGLE_KEY_COL_C - 1, 2)):
            if len(row) > key_idx and isinstance(row[key_idx], str) and row[key_idx].strip():
                values[row[key_idx].strip()] = row[value_idx] if len(row) > value_idx else None
    return Student(row=0, no=1, class_name="", number="", values=values)
